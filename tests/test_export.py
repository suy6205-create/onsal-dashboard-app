import glob
import io
from datetime import date

import openpyxl
import pytest

from src.purchase_order import PO_HEADERS, build_po_xlsx, po_filename

ROWS = [
    {"sku_id": "1000001", "name": "테스트 세럼 30ml 1개", "qty": 20, "prior": 0, "stock": 2, "avg7": 0.286},
    {"sku_id": "1000002", "name": "테스트 로션 150ml 1개", "qty": 10, "prior": 5, "stock": 0, "avg7": 0},
]


def _load(rows=ROWS, vid="A0000001", vname="테스트 주식회사", center="VF107"):
    return openpyxl.load_workbook(io.BytesIO(build_po_xlsx(rows, vid, vname, center))).active


def test_headers_and_sheet_name():
    ws = _load()
    assert ws.title == "발주요청"
    assert [c.value for c in ws[1]] == PO_HEADERS
    assert PO_HEADERS[5] == "기발주량\n(오늘 ~ 차주 토요일까지)"


def test_row_values_formulas_and_formats():
    ws = _load()
    assert [c.value for c in ws[2]][:8] == [1000001, "테스트 세럼 30ml 1개", "A0000001", "테스트 주식회사", 20, 0, 2, 0.3]
    assert ws["A2"].value == 1000001 and isinstance(ws["A2"].value, int)       # SKU ID 는 숫자
    assert ws["I2"].value == '=IFERROR((E2+F2+G2)/H2,"")'                       # 일평균 0 이어도 오류 없음
    assert ws["I3"].value == '=IFERROR((E3+F3+G3)/H3,"")'
    assert ws["J2"].value == "VF107" and ws["K2"].value == "=TODAY()+1" and ws["K2"].number_format == "mm-dd-yy"
    assert "0.0" in ws["H2"].number_format and ws["I2"].number_format.startswith("_-* #,##0_-")
    assert ws["F3"].value == 5 and ws.max_row == 3


def test_header_style_matches_template():
    ws = _load()
    green = {c.coordinate for c in ws[1] if c.fill.fill_type == "solid" and c.fill.fgColor.rgb == "FF92D050"}
    assert green == {"E1", "F1", "K1"}
    assert ws["F1"].alignment.wrap_text and ws["F1"].alignment.horizontal == "center"
    assert ws.row_dimensions[1].height == 24
    assert ws["A1"].font.name == "맑은 고딕" and ws["A1"].font.sz == 9 and ws["B2"].border.left.style == "thin"
    assert ws.column_dimensions["B"].width == 40.375 and ws.column_dimensions["F"].width == 19.375


def test_filename_rules():
    assert po_filename("오아 주식회사", "VF107", date(2026, 10, 19)) == "261019_매뉴얼 발주요청_오아_VF.xlsx"
    assert po_filename("오아 주식회사", "VF107", date(2026, 10, 12), extra=True) == "261012_매뉴얼 발주요청_오아_VF_추가.xlsx"
    assert po_filename("", "", date(2026, 10, 19)) == "261019_매뉴얼 발주요청.xlsx"


REAL = glob.glob("C:/Users/OA/Desktop/로켓용/입고 요청서/261019*.xlsx")


@pytest.mark.skipif(not REAL, reason="사용자의 실제 발주서가 이 PC에 있을 때만 대조")
def test_matches_real_template_cell_by_cell():
    real = openpyxl.load_workbook(REAL[0]).active
    rows = [{"sku_id": r[0], "name": r[1], "qty": r[4], "prior": r[5], "stock": r[6], "avg7": r[7]}
            for r in real.iter_rows(min_row=2, values_only=True)]
    mine = _load(rows, real["C2"].value, real["D2"].value, real["J2"].value)
    for r in range(1, len(rows) + 2):
        for c in range(1, 12):
            a, b = real.cell(r, c), mine.cell(r, c)
            if not (c == 9 and r > 1):                    # I열: 0 나누기 방지로 IFERROR 를 씌운 것만 의도적 차이
                assert a.value == b.value, a.coordinate
            assert a.number_format == b.number_format, a.coordinate
            assert (a.font.name, a.font.sz) == (b.font.name, b.font.sz), a.coordinate
            assert a.border.left.style == b.border.left.style, a.coordinate
    for k in "ABDEFGHIK":
        assert real.column_dimensions[k].width == mine.column_dimensions[k].width
