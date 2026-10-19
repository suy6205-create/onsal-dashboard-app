"""쿠팡 로켓 '매뉴얼 발주요청' 발주서(.xlsx) 생성. 판매·재고 화면에서 사용."""
from __future__ import annotations

from io import BytesIO

PO_HEADERS = ["SKU ID", "스큐 이름", "VENDOR ID", "VENDOR NAME", "발주 요청량",
              "기발주량\n(오늘 ~ 차주 토요일까지)", "쿠팡 현재고", "일평균판매량", "재고일수", "납품센터", "납품일자"]
_GREEN = "FF92D050"
_ACC0 = r'_-* #,##0_-;\-* #,##0_-;_-* "-"_-;_-@_-'
_ACC1 = r'_-* #,##0.0_-;\-* #,##0.0_-;_-* "-"_-;_-@_-'
_PO_WIDTHS = {"A": 13.5, "B": 40.375, "D": 12.5, "E": 12.5, "F": 19.375, "G": 9.5, "H": 10.5, "I": 12.0, "K": 11.125}


def po_filename(vendor_name: str, center: str, today=None, extra: bool = False) -> str:
    """261019_매뉴얼 발주요청_오아_VF.xlsx  (같은 날 추가 발주면 ..._VF_추가.xlsx)"""
    import re
    from datetime import date
    today = today or date.today()
    short = (vendor_name or "").split()[0] if (vendor_name or "").strip() else ""
    m = re.match(r"[A-Za-z]+", center or "")
    parts = [f"{today:%y%m%d}_매뉴얼 발주요청"] + [p for p in (short, m.group(0) if m else "") if p]
    return "_".join(parts) + ("_추가" if extra else "") + ".xlsx"


def build_po_xlsx(rows: list[dict], vendor_id: str, vendor_name: str, center: str = "VF107") -> bytes:
    """사용자가 매번 손으로 만들던 발주서와 같은 서식의 엑셀(.xlsx) 바이트를 만든다.

    rows 의 각 항목: sku_id, name(SKU 명), qty(발주 요청량), prior(기발주량), stock(쿠팡 현재고), avg7(일평균판매량)
    - 재고일수(I)는 수식 =(E+F+G)/H (일평균이 0 이면 빈칸), 납품일자(K)는 =TODAY()+1
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "발주요청"
    font = Font(name="맑은 고딕", size=9)
    side = Side(style="thin")
    border = Border(left=side, right=side, top=side, bottom=side)
    green = PatternFill("solid", fgColor=_GREEN)
    white = PatternFill("solid", fgColor="FFFFFFFF")

    for col, h in enumerate(PO_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.border = font, border
        c.alignment = Alignment(vertical="center", horizontal="center" if col == 6 else None, wrap_text=(col == 6) or None)
        if col in (5, 6, 11):
            c.fill = green
        elif col in (7, 10):
            c.fill = white
    ws.row_dimensions[1].height = 24

    for i, r in enumerate(rows, start=2):
        try:
            sku = int(str(r["sku_id"]).strip())
        except ValueError:
            sku = r["sku_id"]
        vals = [sku, r["name"], vendor_id, vendor_name, int(r["qty"]), int(r.get("prior", 0) or 0), int(r["stock"]),
                round(float(r["avg7"]), 1), f'=IFERROR((E{i}+F{i}+G{i})/H{i},"")', center, "=TODAY()+1"]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font, c.border = font, border
        ws.cell(row=i, column=8).number_format = _ACC1
        ws.cell(row=i, column=9).number_format = _ACC0
        ws.cell(row=i, column=11).number_format = "mm-dd-yy"

    for letter, w in _PO_WIDTHS.items():
        ws.column_dimensions[letter].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
